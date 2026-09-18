"""
E-Studio Desktop Application
============================
Standalone Python Desktop GUI for Label Template Design, Tier Price Resolution,
Data Binding (Excel / CSV), and Sheet Imposition / Vector PDF Generation.

Dependencies:
    pip install PySide6 reportlab python-barcode openpyxl qrcode[pil] pillow

Author: E-Studio Architecture Team
"""

import sys
import os
import json
import math
import io
import csv
from typing import Dict, Any, List, Optional, Tuple

# PySide6 GUI imports
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QTabWidget, QLabel, QPushButton, QFileDialog, QMessageBox,
    QTableWidget, QTableWidgetItem, QHeaderView, QSpinBox, QDoubleSpinBox,
    QComboBox, QCheckBox, QGroupBox, QFormLayout, QLineEdit, QTextEdit,
    QGraphicsView, QGraphicsScene, QGraphicsRectItem, QGraphicsTextItem,
    QGraphicsLineItem, QGraphicsEllipseItem, QGraphicsItem, QSplitter,
    QStatusBar, QToolBar, QProgressBar, QDialog
)
from PySide6.QtGui import (
    QColor, QPen, QBrush, QFont, QPixmap, QImage, QPainter, QAction, QIcon
)
from PySide6.QtCore import Qt, QRectF, QPointF, Signal

# Excel & Data imports
import openpyxl

# ReportLab PDF Generation
from reportlab.lib import pagesizes
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.graphics.shapes import Drawing, Rect, String, Line, Group
from reportlab.graphics.barcode import eanbc, qr
from reportlab.graphics import renderPDF

# Barcode generation for local PIL rendering
import barcode
from barcode.writer import ImageWriter
import qrcode
from PIL import Image

# -----------------------------------------------------------------------------
# 1. DOMAIN & IMPOSITION ENGINES
# -----------------------------------------------------------------------------

PAPER_SIZES_MM = {
    "A4": (210.0, 297.0),
    "A3": (297.0, 420.0),
    "Letter": (215.9, 279.4),
}

DEFAULT_TEMPLATE = {
    "schema_version": 1,
    "name": "Gabarit Standard (100 x 50 mm)",
    "width_mm": 100.0,
    "height_mm": 50.0,
    "inner_margins_mm": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0},
    "outer_margins_mm": {"top": 1.0, "bottom": 1.0, "left": 1.0, "right": 1.0},
    "bg_color": "#FFFFFF",
    "items": [
        {
            "id": "store_name",
            "type": "text",
            "x_mm": 4.0,
            "y_mm": 3.0,
            "w_mm": 45.0,
            "h_mm": 6.0,
            "text": "SUPERMARCHE CENTRAL",
            "binding_key": "STORE_NAME",
            "font_family": "Helvetica",
            "font_size_pt": 8.0,
            "font_weight": "bold",
            "text_color": "#1e293b",
            "alignment": "left"
        },
        {
            "id": "item_category",
            "type": "text",
            "x_mm": 50.0,
            "y_mm": 3.0,
            "w_mm": 46.0,
            "h_mm": 6.0,
            "text": "EPICERIE",
            "binding_key": "CATEGORY_NAME",
            "font_family": "Helvetica",
            "font_size_pt": 8.0,
            "font_weight": "normal",
            "text_color": "#64748b",
            "alignment": "right"
        },
        {
            "id": "div_line",
            "type": "line",
            "x_mm": 4.0,
            "y_mm": 10.0,
            "w_mm": 92.0,
            "h_mm": 1.0,
            "color": "#cbd5e1",
            "thickness": 1.0
        },
        {
            "id": "item_name",
            "type": "text",
            "x_mm": 4.0,
            "y_mm": 12.0,
            "w_mm": 58.0,
            "h_mm": 16.0,
            "text": "Nom du Produit Exemple",
            "binding_key": "ITEMNAME",
            "font_family": "Helvetica",
            "font_size_pt": 12.0,
            "font_weight": "bold",
            "text_color": "#0f172a",
            "alignment": "left"
        },
        {
            "id": "price_bg",
            "type": "shape",
            "x_mm": 62.0,
            "y_mm": 12.0,
            "w_mm": 34.0,
            "h_mm": 18.0,
            "fill_color": "#eff6ff",
            "border_color": "#3b82f6",
            "border_width": 1.0,
            "corner_radius": 4.0
        },
        {
            "id": "item_price",
            "type": "text",
            "x_mm": 64.0,
            "y_mm": 15.0,
            "w_mm": 30.0,
            "h_mm": 12.0,
            "text": "2 500 FCFA",
            "binding_key": "SELLING_PRICE",
            "font_family": "Helvetica",
            "font_size_pt": 14.0,
            "font_weight": "bold",
            "text_color": "#1d4ed8",
            "alignment": "center"
        },
        {
            "id": "barcode_elem",
            "type": "barcode",
            "x_mm": 4.0,
            "y_mm": 30.0,
            "w_mm": 55.0,
            "h_mm": 15.0,
            "code": "3017620422003",
            "barcode_type": "ean13",
            "binding_key": "BARCODE",
            "show_text": True,
            "bar_color": "#000000"
        },
        {
            "id": "tier_price_elem",
            "type": "tier_price",
            "x_mm": 62.0,
            "y_mm": 32.0,
            "w_mm": 34.0,
            "h_mm": 14.0,
            "primary_tier": 1,
            "prefix_text": "Des",
            "unit_label": "FCFA",
            "strict_required": False,
            "fallback_to_base_price": True
        }
    ]
}

class TierEngine:
    @staticmethod
    def resolve(item_prop: Dict[str, Any], record: Dict[str, Any]) -> Dict[str, Any]:
        primary_tier = item_prop.get("primary_tier", 1)
        prefix_text = item_prop.get("prefix_text", "Des")
        unit_label = item_prop.get("unit_label", "FCFA")
        strict_required = item_prop.get("strict_required", False)
        fallback = item_prop.get("fallback_to_base_price", True)

        tiers = record.get("TIERS") or []
        tier = None
        target_idx = primary_tier - 1
        if 0 <= target_idx < len(tiers):
            tier = tiers[target_idx]

        if tier:
            qty = tier.get("qty", 1)
            uprice = tier.get("unit_price", 0)
            return {
                "text_qty": f"{prefix_text} {qty} pces",
                "formatted_price": f"{uprice:,.0f} {unit_label}".replace(",", " "),
                "is_fallback": False,
                "error": None
            }

        if strict_required:
            return {
                "text_qty": "[Palier Manquant]",
                "formatted_price": "N/A",
                "is_fallback": False,
                "error": f"Palier #{primary_tier} non trouve"
            }

        if fallback:
            base_p = record.get("SELLING_PRICE", 0)
            try:
                base_val = float(str(base_p).replace(" ", "").replace(",", "."))
            except:
                base_val = 0
            return {
                "text_qty": "Prix Standard",
                "formatted_price": f"{base_val:,.0f} {unit_label}".replace(",", " "),
                "is_fallback": True,
                "error": None
            }

        return {
            "text_qty": "",
            "formatted_price": "",
            "is_fallback": False,
            "error": "Aucun prix"
        }

class ImpositionEngine:
    @staticmethod
    def calculate(
        label_w_mm: float,
        label_h_mm: float,
        page_size_name: str = "A4",
        orientation: str = "portrait",
        gap_h_mm: float = 2.0,
        gap_v_mm: float = 2.0,
        margin_h_mm: float = 5.0,
        margin_v_mm: float = 5.0
    ) -> Dict[str, Any]:
        p_w, p_h = PAPER_SIZES_MM.get(page_size_name, (210.0, 297.0))
        if orientation == "landscape":
            p_w, p_h = max(p_w, p_h), min(p_w, p_h)
        else:
            p_w, p_h = min(p_w, p_h), max(p_w, p_h)

        usable_w = p_w - 2 * margin_h_mm
        usable_h = p_h - 2 * margin_v_mm

        cols = max(1, math.floor((usable_w + gap_h_mm) / (label_w_mm + gap_h_mm)))
        rows = max(1, math.floor((usable_h + gap_v_mm) / (label_h_mm + gap_v_mm)))

        grid_w = cols * label_w_mm + (cols - 1) * gap_h_mm
        grid_h = rows * label_h_mm + (rows - 1) * gap_v_mm

        offset_x = (p_w - grid_w) / 2.0
        offset_y = (p_h - grid_h) / 2.0

        return {
            "page_w_mm": p_w,
            "page_h_mm": p_h,
            "cols": cols,
            "rows": rows,
            "total_per_page": cols * rows,
            "offset_x_mm": offset_x,
            "offset_y_mm": offset_y,
            "grid_w_mm": grid_w,
            "grid_h_mm": grid_h
        }

# -----------------------------------------------------------------------------
# 2. VECTOR PDF EXPORTER (ReportLab)
# -----------------------------------------------------------------------------

def hex_to_rgb(hex_str: str) -> Tuple[float, float, float]:
    if not hex_str or not hex_str.startswith("#"):
        return (0.0, 0.0, 0.0)
    h = hex_str.lstrip("#")
    if len(h) == 3:
        h = "".join([c*2 for c in h])
    if len(h) >= 6:
        r = int(h[0:2], 16) / 255.0
        g = int(h[2:4], 16) / 255.0
        b = int(h[4:6], 16) / 255.0
        return (r, g, b)
    return (0.0, 0.0, 0.0)

def draw_single_label_reportlab(
    pdf_c: canvas.Canvas,
    origin_x_pt: float,
    origin_y_pt: float,
    template: Dict[str, Any],
    record: Dict[str, Any]
):
    """Draws a label at origin_x_pt, origin_y_pt (bottom-left coordinate in points)."""
    lbl_w_pt = template["width_mm"] * mm
    lbl_h_pt = template["height_mm"] * mm

    pdf_c.saveState()
    pdf_c.translate(origin_x_pt, origin_y_pt)

    # Background
    bg_color = template.get("bg_color", "#FFFFFF")
    r, g, b = hex_to_rgb(bg_color)
    pdf_c.setFillColorRGB(r, g, b)
    pdf_c.rect(0, 0, lbl_w_pt, lbl_h_pt, fill=1, stroke=0)

    # Optional border / inner margin preview
    pdf_c.setStrokeColorRGB(0.85, 0.85, 0.85)
    pdf_c.setLineWidth(0.3)
    pdf_c.rect(0, 0, lbl_w_pt, lbl_h_pt, fill=0, stroke=1)

    # Render items
    items = template.get("items", [])
    for item in items:
        itype = item.get("type")
        x = item.get("x_mm", 0) * mm
        # In ReportLab, y goes upwards. Let's convert top-down mm to bottom-up pt:
        # y_from_top = item['y_mm'] * mm
        # y_reportlab = lbl_h_pt - (y_from_top + h_mm * mm)
        w = item.get("w_mm", 10) * mm
        h = item.get("h_mm", 5) * mm
        y = lbl_h_pt - (item.get("y_mm", 0) * mm + h)

        if itype == "shape":
            fill = item.get("fill_color", "#ffffff")
            border = item.get("border_color", "#000000")
            rf, gf, bf = hex_to_rgb(fill)
            rb, gb, bb = hex_to_rgb(border)
            pdf_c.setFillColorRGB(rf, gf, bf)
            pdf_c.setStrokeColorRGB(rb, gb, bb)
            pdf_c.setLineWidth(item.get("border_width", 1.0))
            corner = item.get("corner_radius", 0.0) * mm
            pdf_c.roundRect(x, y, w, h, corner, fill=1, stroke=1)

        elif itype == "line":
            col = item.get("color", "#000000")
            rl, gl, bl = hex_to_rgb(col)
            pdf_c.setStrokeColorRGB(rl, gl, bl)
            pdf_c.setLineWidth(item.get("thickness", 1.0))
            pdf_c.line(x, y + h / 2.0, x + w, y + h / 2.0)

        elif itype == "text":
            b_key = item.get("binding_key")
            val_text = str(record.get(b_key) if b_key and b_key in record else item.get("text", ""))

            # Currency format if SELLING_PRICE
            if b_key == "SELLING_PRICE":
                try:
                    numeric_val = float(str(val_text).replace(" ", "").replace(",", "."))
                    val_text = f"{numeric_val:,.0f} FCFA".replace(",", " ")
                except:
                    pass

            txt_col = item.get("text_color", "#000000")
            rt, gt, bt = hex_to_rgb(txt_col)
            pdf_c.setFillColorRGB(rt, gt, bt)

            font_weight = item.get("font_weight", "normal")
            font_name = "Helvetica-Bold" if font_weight in ("bold", "600", "800") else "Helvetica"
            font_sz = item.get("font_size_pt", 10.0)
            pdf_c.setFont(font_name, font_sz)

            align = item.get("alignment", "left")
            text_baseline = y + (h / 2.0) - (font_sz / 3.0)

            if align == "center":
                pdf_c.drawCentredString(x + w / 2.0, text_baseline, val_text)
            elif align == "right":
                pdf_c.drawRightString(x + w, text_baseline, val_text)
            else:
                pdf_c.drawString(x, text_baseline, val_text)

        elif itype == "tier_price":
            tier_info = TierEngine.resolve(item, record)
            # Box background
            pdf_c.setFillColorRGB(0.97, 0.98, 1.0)
            pdf_c.setStrokeColorRGB(0.7, 0.8, 0.95)
            pdf_c.setLineWidth(0.8)
            pdf_c.roundRect(x, y, w, h, 2*mm, fill=1, stroke=1)

            # Draw Qty prefix
            pdf_c.setFillColorRGB(0.3, 0.4, 0.5)
            pdf_c.setFont("Helvetica", 7.0)
            pdf_c.drawCentredString(x + w / 2.0, y + h - 3.5*mm, tier_info["text_qty"])

            # Draw Tier Price
            pdf_c.setFillColorRGB(0.1, 0.3, 0.8)
            pdf_c.setFont("Helvetica-Bold", 9.5)
            pdf_c.drawCentredString(x + w / 2.0, y + 2.0*mm, tier_info["formatted_price"])

        elif itype == "barcode":
            code_val = str(record.get("BARCODE") or record.get("PRODUCT_SCAN") or item.get("code", "0000000000000"))
            b_type = item.get("barcode_type", "ean13")
            try:
                if b_type == "ean13" and len(code_val) == 13 and code_val.isdigit():
                    bc = eanbc.Ean13BarcodeWidget(code_val)
                    bc.barWidth = 0.9
                    bc.barHeight = h - (4*mm if item.get("show_text", True) else 0)
                    bc.drawOn(pdf_c, x + 2*mm, y)
                else:
                    # Generic visual fallback barcode representation
                    pdf_c.setFillColorRGB(0, 0, 0)
                    bars = 30
                    bw = w / (bars * 1.5)
                    for bi in range(bars):
                        if (bi % 3) != 0:
                            pdf_c.rect(x + bi * bw * 1.5, y + 3*mm, bw, h - 3*mm, fill=1, stroke=0)
                    if item.get("show_text", True):
                        pdf_c.setFont("Courier", 7.0)
                        pdf_c.drawCentredString(x + w / 2.0, y, code_val)
            except Exception as e:
                pdf_c.setFont("Helvetica", 8.0)
                pdf_c.setFillColorRGB(0.8, 0, 0)
                pdf_c.drawString(x, y + h/2.0, f"Code: {code_val}")

        elif itype == "qrcode":
            qr_text = str(record.get("BARCODE") or item.get("content", "http://e-studio.local"))
            try:
                qr_code = qr.QrCodeWidget(qr_text)
                qr_code.barWidth = min(w, h)
                qr_code.barHeight = min(w, h)
                qr_code.drawOn(pdf_c, x, y)
            except:
                pass

    pdf_c.restoreState()


def export_imposed_pdf(
    out_path: str,
    template: Dict[str, Any],
    records: List[Dict[str, Any]],
    page_size_name: str = "A4",
    orientation: str = "portrait",
    draw_crop_marks: bool = True
):
    dim1, dim2 = (210.0*mm, 297.0*mm) if page_size_name == "A4" else (
        (297.0*mm, 420.0*mm) if page_size_name == "A3" else (215.9*mm, 279.4*mm)
    )
    pw = min(dim1, dim2) if orientation == "portrait" else max(dim1, dim2)
    ph = max(dim1, dim2) if orientation == "portrait" else min(dim1, dim2)

    imposition = ImpositionEngine.calculate(
        template["width_mm"],
        template["height_mm"],
        page_size_name=page_size_name,
        orientation=orientation
    )

    pdf_c = canvas.Canvas(out_path, pagesize=(pw, ph))
    cols = imposition["cols"]
    rows = imposition["rows"]
    per_page = cols * rows

    lbl_w_pt = template["width_mm"] * mm
    lbl_h_pt = template["height_mm"] * mm
    gap_pt = 2.0 * mm

    offset_x_pt = imposition["offset_x_mm"] * mm
    offset_y_pt = imposition["offset_y_mm"] * mm

    rec_idx = 0
    total_records = len(records)

    while rec_idx < total_records:
        for r_i in range(rows):
            for c_i in range(cols):
                if rec_idx >= total_records:
                    break
                rec = records[rec_idx]

                # Calculate slot coordinate
                # row index 0 is at top
                slot_x = offset_x_pt + c_i * (lbl_w_pt + gap_pt)
                slot_y = ph - (offset_y_pt + (r_i + 1) * lbl_h_pt + r_i * gap_pt)

                # Draw label
                draw_single_label_reportlab(pdf_c, slot_x, slot_y, template, rec)

                # Optional crop marks
                if draw_crop_marks:
                    pdf_c.setStrokeColorRGB(0.7, 0.7, 0.7)
                    pdf_c.setLineWidth(0.4)
                    mark_len = 3 * mm
                    # Top-left mark
                    pdf_c.line(slot_x - 1*mm, slot_y + lbl_h_pt, slot_x - 1*mm - mark_len, slot_y + lbl_h_pt)
                    pdf_c.line(slot_x, slot_y + lbl_h_pt + 1*mm, slot_x, slot_y + lbl_h_pt + 1*mm + mark_len)
                    # Bottom-right mark
                    pdf_c.line(slot_x + lbl_w_pt + 1*mm, slot_y, slot_x + lbl_w_pt + 1*mm + mark_len, slot_y)
                    pdf_c.line(slot_x + lbl_w_pt, slot_y - 1*mm, slot_x + lbl_w_pt, slot_y - 1*mm - mark_len)

                rec_idx += 1

        pdf_c.showPage()

    pdf_c.save()


# -----------------------------------------------------------------------------
# 3. INTERACTIVE PYSIDE6 CANVAS ITEM
# -----------------------------------------------------------------------------

SCALE_FACTOR = 3.78  # 1 mm ≈ 3.78 px at 96 DPI

class VisualLabelItem(QGraphicsRectItem):
    def __init__(self, item_prop: Dict[str, Any], parent=None):
        super().__init__(parent)
        self.item_prop = item_prop
        self.setFlags(
            QGraphicsItem.ItemIsMovable |
            QGraphicsItem.ItemIsSelectable |
            QGraphicsItem.ItemSendsGeometryChanges
        )
        self.update_geometry_from_prop()

    def update_geometry_from_prop(self):
        x = self.item_prop.get("x_mm", 0) * SCALE_FACTOR
        y = self.item_prop.get("y_mm", 0) * SCALE_FACTOR
        w = self.item_prop.get("w_mm", 20) * SCALE_FACTOR
        h = self.item_prop.get("h_mm", 10) * SCALE_FACTOR
        self.setRect(0, 0, w, h)
        self.setPos(x, y)

    def paint(self, painter: QPainter, option, widget=None):
        rect = self.rect()
        itype = self.item_prop.get("type", "text")

        if itype == "shape":
            fc = QColor(self.item_prop.get("fill_color", "#ffffff"))
            bc = QColor(self.item_prop.get("border_color", "#000000"))
            painter.setBrush(QBrush(fc))
            painter.setPen(QPen(bc, self.item_prop.get("border_width", 1.0)))
            cr = self.item_prop.get("corner_radius", 0.0) * SCALE_FACTOR
            painter.drawRoundedRect(rect, cr, cr)

        elif itype == "line":
            lc = QColor(self.item_prop.get("color", "#cbd5e1"))
            painter.setPen(QPen(lc, self.item_prop.get("thickness", 1.0) * SCALE_FACTOR / 2.0))
            mid_y = rect.top() + rect.height() / 2.0
            painter.drawLine(rect.left(), mid_y, rect.right(), mid_y)

        elif itype in ("text", "tier_price"):
            is_tier = (itype == "tier_price")
            if is_tier:
                painter.setBrush(QBrush(QColor("#eff6ff")))
                painter.setPen(QPen(QColor("#3b82f6"), 1))
                painter.drawRoundedRect(rect, 4, 4)
                painter.setPen(QColor("#1d4ed8"))
                painter.setFont(QFont("Helvetica", 9, QFont.Bold))
                pfx = self.item_prop.get("prefix_text", "Des")
                lbl = self.item_prop.get("unit_label", "FCFA")
                painter.drawText(rect, Qt.AlignCenter, f"{pfx} 5 pces\n2 100 {lbl}")
            else:
                tc = QColor(self.item_prop.get("text_color", "#0f172a"))
                painter.setPen(tc)
                fsz = self.item_prop.get("font_size_pt", 10.0) * 0.9
                weight = QFont.Bold if self.item_prop.get("font_weight") in ("bold", "600") else QFont.Normal
                painter.setFont(QFont("Helvetica", int(fsz), weight))
                txt = self.item_prop.get("text", "Texte")
                bkey = self.item_prop.get("binding_key")
                if bkey:
                    txt = f"[{bkey}]"
                painter.drawText(rect, Qt.AlignVCenter | Qt.AlignLeft, txt)

        elif itype == "barcode":
            painter.setBrush(QBrush(QColor("#ffffff")))
            painter.setPen(QPen(QColor("#cbd5e1"), 1))
            painter.drawRect(rect)
            painter.setPen(QColor("#000000"))
            painter.setFont(QFont("Courier", 7))
            painter.drawText(rect, Qt.AlignCenter, f"||||| EAN13: {self.item_prop.get('code', '3017620422003')} |||||")

        if self.isSelected():
            painter.setBrush(Qt.NoBrush)
            painter.setPen(QPen(QColor("#2563eb"), 2, Qt.DashLine))
            painter.drawRect(rect)

    def itemChange(self, change, value):
        if change == QGraphicsItem.ItemPositionChange and self.scene():
            new_pos = value
            self.item_prop["x_mm"] = round(new_pos.x() / SCALE_FACTOR, 1)
            self.item_prop["y_mm"] = round(new_pos.y() / SCALE_FACTOR, 1)
        return super().itemChange(change, value)


# -----------------------------------------------------------------------------
# 4. MAIN DESKTOP APPLICATION WINDOW (PySide6)
# -----------------------------------------------------------------------------

class EStudioMainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("E-Studio - Etiquettes & Imposition (PySide6 / ReportLab)")
        self.resize(1280, 850)

        self.template = json.loads(json.dumps(DEFAULT_TEMPLATE))
        self.records: List[Dict[str, Any]] = [
            {
                "id": "1",
                "STORE_NAME": "SUPERMARCHE EXPRESS",
                "CATEGORY_NAME": "EPICERIE",
                "ITEMNAME": "Cafe Moulu Pur Arabica 250g",
                "SELLING_PRICE": 2500,
                "BARCODE": "3017620422003",
                "TIERS": [{"qty": 5, "unit_price": 2250}, {"qty": 10, "unit_price": 2000}]
            },
            {
                "id": "2",
                "STORE_NAME": "SUPERMARCHE EXPRESS",
                "CATEGORY_NAME": "BOISSONS",
                "ITEMNAME": "Jus d'Orange Bio 1L",
                "SELLING_PRICE": 1800,
                "BARCODE": "5449000000996",
                "TIERS": [{"qty": 6, "unit_price": 1600}]
            },
            {
                "id": "3",
                "STORE_NAME": "SUPERMARCHE EXPRESS",
                "CATEGORY_NAME": "PRODUITS FRAIS",
                "ITEMNAME": "Lait Demi-Ecreme 1L",
                "SELLING_PRICE": 950,
                "BARCODE": "3250390111111",
                "TIERS": []
            }
        ]

        self.init_ui()

    def init_ui(self):
        self.tabs = QTabWidget()
        self.setCentralWidget(self.tabs)

        # Tab 1: Visual Template Designer
        self.designer_tab = QWidget()
        self.init_designer_tab()
        self.tabs.addTab(self.designer_tab, "1. Editeur de Gabarit (Canvas)")

        # Tab 2: Batch Data & Records
        self.data_tab = QWidget()
        self.init_data_tab()
        self.tabs.addTab(self.data_tab, "2. Donnees & Paliers (Excel / CSV)")

        # Tab 3: Sheet Imposition & PDF Export
        self.imposition_tab = QWidget()
        self.init_imposition_tab()
        self.tabs.addTab(self.imposition_tab, "3. Imposition & Export PDF")

        self.status = QStatusBar()
        self.setStatusBar(self.status)
        self.status.showMessage("Pret - E-Studio Python Desktop")

    # -------------------------------------------------------------------------
    # TAB 1: VISUAL DESIGNER
    # -------------------------------------------------------------------------
    def init_designer_tab(self):
        layout = QHBoxLayout(self.designer_tab)

        # Canvas Graphics View
        canvas_group = QGroupBox("Canvas d'Etiquette Interactif (Milimetrique)")
        canvas_layout = QVBoxLayout(canvas_group)

        self.scene = QGraphicsScene()
        self.view = QGraphicsView(self.scene)
        self.view.setRenderHint(QPainter.Antialiasing)
        self.view.setStyleSheet("background-color: #f1f5f9;")
        canvas_layout.addWidget(self.view)

        # Controls panel
        props_panel = QGroupBox("Proprietes du Gabarit")
        props_layout = QFormLayout(props_panel)
        props_panel.setFixedWidth(320)

        self.spin_width = QDoubleSpinBox()
        self.spin_width.setRange(20, 300)
        self.spin_width.setValue(self.template["width_mm"])
        self.spin_width.valueChanged.connect(self.on_template_dimensions_changed)

        self.spin_height = QDoubleSpinBox()
        self.spin_height.setRange(15, 300)
        self.spin_height.setValue(self.template["height_mm"])
        self.spin_height.valueChanged.connect(self.on_template_dimensions_changed)

        props_layout.addRow("Largeur (mm):", self.spin_width)
        props_layout.addRow("Hauteur (mm):", self.spin_height)

        btn_add_text = QPushButton("+ Ajouter Texte")
        btn_add_text.clicked.connect(self.add_text_item)
        props_layout.addRow(btn_add_text)

        btn_add_barcode = QPushButton("+ Ajouter Code-barres")
        btn_add_barcode.clicked.connect(self.add_barcode_item)
        props_layout.addRow(btn_add_barcode)

        btn_add_tier = QPushButton("+ Ajouter Prix Palier")
        btn_add_tier.clicked.connect(self.add_tier_item)
        props_layout.addRow(btn_add_tier)

        btn_save_tpl = QPushButton("Sauvegarder Gabarit JSON")
        btn_save_tpl.clicked.connect(self.save_template_json)
        props_layout.addRow(btn_save_tpl)

        btn_load_tpl = QPushButton("Ouvrir Gabarit JSON")
        btn_load_tpl.clicked.connect(self.load_template_json)
        props_layout.addRow(btn_load_tpl)

        layout.addWidget(canvas_group, 3)
        layout.addWidget(props_panel, 1)

        self.refresh_canvas()

    def refresh_canvas(self):
        self.scene.clear()
        w_px = self.template["width_mm"] * SCALE_FACTOR
        h_px = self.template["height_mm"] * SCALE_FACTOR

        # Label background
        bg_rect = QGraphicsRectItem(0, 0, w_px, h_px)
        bg_rect.setBrush(QBrush(QColor(self.template.get("bg_color", "#ffffff"))))
        bg_rect.setPen(QPen(QColor("#94a3b8"), 1))
        self.scene.addItem(bg_rect)

        # Add items
        for item in self.template.get("items", []):
            visual = VisualLabelItem(item)
            self.scene.addItem(visual)

        self.scene.setSceneRect(-20, -20, w_px + 40, h_px + 40)

    def on_template_dimensions_changed(self):
        self.template["width_mm"] = self.spin_width.value()
        self.template["height_mm"] = self.spin_height.value()
        self.refresh_canvas()

    def add_text_item(self):
        new_item = {
            "id": f"text_{len(self.template['items'])+1}",
            "type": "text",
            "x_mm": 5.0,
            "y_mm": 5.0,
            "w_mm": 40.0,
            "h_mm": 8.0,
            "text": "Nouveau Texte",
            "binding_key": "ITEMNAME",
            "font_family": "Helvetica",
            "font_size_pt": 10.0,
            "font_weight": "normal",
            "text_color": "#000000",
            "alignment": "left"
        }
        self.template["items"].append(new_item)
        self.refresh_canvas()

    def add_barcode_item(self):
        new_item = {
            "id": f"barcode_{len(self.template['items'])+1}",
            "type": "barcode",
            "x_mm": 5.0,
            "y_mm": 18.0,
            "w_mm": 50.0,
            "h_mm": 15.0,
            "code": "3017620422003",
            "barcode_type": "ean13",
            "binding_key": "BARCODE",
            "show_text": True,
            "bar_color": "#000000"
        }
        self.template["items"].append(new_item)
        self.refresh_canvas()

    def add_tier_item(self):
        new_item = {
            "id": f"tier_{len(self.template['items'])+1}",
            "type": "tier_price",
            "x_mm": 60.0,
            "y_mm": 18.0,
            "w_mm": 35.0,
            "h_mm": 14.0,
            "primary_tier": 1,
            "prefix_text": "Des",
            "unit_label": "FCFA",
            "strict_required": False,
            "fallback_to_base_price": True
        }
        self.template["items"].append(new_item)
        self.refresh_canvas()

    def save_template_json(self):
        path, _ = QFileDialog.getSaveFileName(self, "Enregistrer Gabarit", "gabarit.json", "JSON Files (*.json)")
        if path:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(self.template, f, indent=2, ensure_ascii=False)
            self.status.showMessage(f"Gabarit enregistre sous {path}")

    def load_template_json(self):
        path, _ = QFileDialog.getOpenFileName(self, "Ouvrir Gabarit", "", "JSON Files (*.json)")
        if path:
            with open(path, "r", encoding="utf-8") as f:
                self.template = json.load(f)
            self.spin_width.setValue(self.template.get("width_mm", 100.0))
            self.spin_height.setValue(self.template.get("height_mm", 50.0))
            self.refresh_canvas()
            self.status.showMessage(f"Gabarit charge: {self.template.get('name')}")

    # -------------------------------------------------------------------------
    # TAB 2: DATA & RECORDS (EXCEL & CSV)
    # -------------------------------------------------------------------------
    def init_data_tab(self):
        layout = QVBoxLayout(self.data_tab)

        top_bar = QHBoxLayout()
        btn_import_excel = QPushButton("Importer Fichier Excel (.xlsx)")
        btn_import_excel.clicked.connect(self.import_excel)
        top_bar.addWidget(btn_import_excel)

        btn_import_csv = QPushButton("Importer Fichier CSV (.csv)")
        btn_import_csv.clicked.connect(self.import_csv)
        top_bar.addWidget(btn_import_csv)

        top_bar.addStretch()
        self.lbl_record_count = QLabel("3 articles charges")
        top_bar.addWidget(self.lbl_record_count)

        layout.addLayout(top_bar)

        self.table = QTableWidget()
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels([
            "Magasin", "Rayon", "Designation", "Prix Vente", "Code-barres", "Paliers (Tiers)"
        ])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        layout.addWidget(self.table)

        self.populate_table()

    def populate_table(self):
        self.table.setRowCount(len(self.records))
        for row, rec in enumerate(self.records):
            self.table.setItem(row, 0, QTableWidgetItem(str(rec.get("STORE_NAME", ""))))
            self.table.setItem(row, 1, QTableWidgetItem(str(rec.get("CATEGORY_NAME", ""))))
            self.table.setItem(row, 2, QTableWidgetItem(str(rec.get("ITEMNAME", ""))))
            self.table.setItem(row, 3, QTableWidgetItem(f"{rec.get('SELLING_PRICE', 0)} FCFA"))
            self.table.setItem(row, 4, QTableWidgetItem(str(rec.get("BARCODE", ""))))

            tiers = rec.get("TIERS", [])
            tier_str = ", ".join([f">={t['qty']}pcs: {t['unit_price']}F" for t in tiers]) if tiers else "Aucun"
            self.table.setItem(row, 5, QTableWidgetItem(tier_str))

        self.lbl_record_count.setText(f"{len(self.records)} articles charges")

    def import_excel(self):
        path, _ = QFileDialog.getOpenFileName(self, "Choisir fichier Excel", "", "Excel Files (*.xlsx *.xls)")
        if not path:
            return

        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                QMessageBox.warning(self, "Erreur", "Le fichier Excel ne contient pas de donnees.")
                return

            headers = [str(h).strip().upper() if h else f"COL_{idx}" for idx, h in enumerate(rows[0])]
            imported: List[Dict[str, Any]] = []

            for row_idx, r in enumerate(rows[1:], start=2):
                if not any(r):
                    continue
                d = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}

                # Parse basic fields
                rec = {
                    "id": str(row_idx),
                    "STORE_NAME": d.get("STORE_NAME") or d.get("MAGASIN") or "SUPERMARCHE",
                    "CATEGORY_NAME": d.get("CATEGORY_NAME") or d.get("RAYON") or "DIVERS",
                    "ITEMNAME": d.get("ITEMNAME") or d.get("DESIGNATION") or d.get("ARTICLE") or "Sans Nom",
                    "SELLING_PRICE": d.get("SELLING_PRICE") or d.get("PRIX") or d.get("PRIX_VENTE") or 0,
                    "BARCODE": str(d.get("BARCODE") or d.get("EAN") or d.get("CB") or ""),
                    "TIERS": []
                }

                # Scan for tier columns e.g. TIER_1_QTY, TIER_1_PRICE
                for t_idx in range(1, 10):
                    q_key = f"TIER_{t_idx}_QTY"
                    p_key = f"TIER_{t_idx}_PRICE"
                    if q_key in d and p_key in d and d[q_key] and d[p_key]:
                        try:
                            rec["TIERS"].append({
                                "qty": int(d[q_key]),
                                "unit_price": float(d[p_key])
                            })
                        except:
                            pass

                imported.append(rec)

            self.records = imported
            self.populate_table()
            QMessageBox.information(self, "Import Reussi", f"{len(imported)} articles importes depuis Excel !")
        except Exception as e:
            QMessageBox.critical(self, "Erreur Import", f"Echec de lecture du fichier Excel: {str(e)}")

    def import_csv(self):
        path, _ = QFileDialog.getOpenFileName(self, "Choisir fichier CSV", "", "CSV Files (*.csv)")
        if not path:
            return

        try:
            with open(path, "r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                imported = []
                for idx, row in enumerate(reader, start=1):
                    rec = {
                        "id": str(idx),
                        "STORE_NAME": row.get("STORE_NAME", "SUPERMARCHE"),
                        "CATEGORY_NAME": row.get("CATEGORY_NAME", "DIVERS"),
                        "ITEMNAME": row.get("ITEMNAME", "Article"),
                        "SELLING_PRICE": row.get("SELLING_PRICE", 0),
                        "BARCODE": row.get("BARCODE", ""),
                        "TIERS": []
                    }
                    imported.append(rec)
            self.records = imported
            self.populate_table()
            QMessageBox.information(self, "Import Reussi", f"{len(imported)} articles importes depuis CSV !")
        except Exception as e:
            QMessageBox.critical(self, "Erreur Import", f"Echec de lecture du fichier CSV: {str(e)}")

    # -------------------------------------------------------------------------
    # TAB 3: IMPOSITION & PDF EXPORT
    # -------------------------------------------------------------------------
    def init_imposition_tab(self):
        layout = QHBoxLayout(self.imposition_tab)

        # Imposition Settings Form
        settings_group = QGroupBox("Configuration de la Planche d'Impression")
        form = QFormLayout(settings_group)
        settings_group.setFixedWidth(380)

        self.combo_paper = QComboBox()
        self.combo_paper.addItems(["A4", "A3", "Letter"])
        self.combo_paper.currentTextChanged.connect(self.update_imposition_summary)
        form.addRow("Format Papier:", self.combo_paper)

        self.combo_orientation = QComboBox()
        self.combo_orientation.addItems(["portrait", "landscape"])
        self.combo_orientation.currentTextChanged.connect(self.update_imposition_summary)
        form.addRow("Orientation:", self.combo_orientation)

        self.check_crop_marks = QCheckBox("Tracer les traits de coupe (Crop Marks)")
        self.check_crop_marks.setChecked(True)
        form.addRow(self.check_crop_marks)

        # Summary box
        self.lbl_imposition_calc = QLabel("")
        self.lbl_imposition_calc.setStyleSheet("background-color: #f8fafc; padding: 12px; border: 1px solid #e2e8f0; border-radius: 6px;")
        form.addRow("Calcul d'Imposition:", self.lbl_imposition_calc)

        btn_generate_pdf = QPushButton("Generer le Fichier PDF Vectoriel")
        btn_generate_pdf.setStyleSheet("background-color: #2563eb; color: white; font-weight: bold; padding: 10px; border-radius: 6px;")
        btn_generate_pdf.clicked.connect(self.generate_pdf)
        form.addRow(btn_generate_pdf)

        layout.addWidget(settings_group)

        # Preview info
        preview_group = QGroupBox("Apercu de la Disposition")
        pv_layout = QVBoxLayout(preview_group)
        self.lbl_visual_preview = QLabel("La disposition automatique optimise les poses pour minimiser les chutes de papier.")
        self.lbl_visual_preview.setAlignment(Qt.AlignCenter)
        self.lbl_visual_preview.setStyleSheet("color: #64748b; font-size: 13px;")
        pv_layout.addWidget(self.lbl_visual_preview)
        layout.addWidget(preview_group)

        self.update_imposition_summary()

    def update_imposition_summary(self):
        paper = self.combo_paper.currentText()
        orientation = self.combo_orientation.currentText()
        imp = ImpositionEngine.calculate(
            self.template["width_mm"],
            self.template["height_mm"],
            page_size_name=paper,
            orientation=orientation
        )

        total_pages = math.ceil(len(self.records) / imp["total_per_page"]) if imp["total_per_page"] > 0 else 1

        summary = (
            f"<b>Grille:</b> {imp['cols']} colonnes x {imp['rows']} lignes<br>"
            f"<b>Etiquettes / Page:</b> {imp['total_per_page']} poses<br>"
            f"<b>Centrage X:</b> {imp['offset_x_mm']:.1f} mm | <b>Y:</b> {imp['offset_y_mm']:.1f} mm<br>"
            f"<b>Total Planches Necessaires:</b> {total_pages} page(s) pour {len(self.records)} articles"
        )
        self.lbl_imposition_calc.setText(summary)

    def generate_pdf(self):
        if not self.records:
            QMessageBox.warning(self, "Aucune donnee", "Veuillez charger au moins un article dans l'onglet Donnees.")
            return

        path, _ = QFileDialog.getSaveFileName(self, "Enregistrer Planche PDF", "planches_etiquettes.pdf", "PDF Files (*.pdf)")
        if not path:
            return

        try:
            paper = self.combo_paper.currentText()
            orientation = self.combo_orientation.currentText()
            crop_marks = self.check_crop_marks.isChecked()

            export_imposed_pdf(
                out_path=path,
                template=self.template,
                records=self.records,
                page_size_name=paper,
                orientation=orientation,
                draw_crop_marks=crop_marks
            )

            QMessageBox.information(self, "Succes", f"Fichier PDF vectoriel genere avec succes !\nEmplacement: {path}")
            self.status.showMessage(f"PDF genere: {path}")
        except Exception as e:
            QMessageBox.critical(self, "Erreur de Generation", f"Echec de generation du PDF: {str(e)}")


# -----------------------------------------------------------------------------
# 5. ENTRY POINT
# -----------------------------------------------------------------------------

def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    win = EStudioMainWindow()
    win.show()
    sys.exit(app.exec())

if __name__ == "__main__":
    main()
