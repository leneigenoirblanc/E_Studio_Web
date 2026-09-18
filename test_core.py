"""Regression checks for the pure-Python E-Studio domain services."""

import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from article_models import PricingInfo, ProductIdent, TaxInfo
from article_types import BulkCaseArticle, WeightedArticle
from imposition import ImpositionCalculator
from template_model import LabelTemplate, Margins
from tier_engine import TierResolver, TierTargetSpec
from binding_profiles import BindingProfile, ColumnMapping
from batch_export import export_rows_as_png, normalize_workbook
from data_normalization import normalize_row
from data_sources import ImportedRow, read_xlsx
from app_settings import AppSettings, BindingObjectMode, GenerationEditPolicy
from tier_grouping import TierGroupingConfig, group_article_rows
from domain_fields import DomainFieldKey, canonical_domain_key
from generation_session import GenerationSession, GenerationState, TemplateSnapshot
from template_library import TemplateLibrary

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover - dependency installation concern
    Workbook = None


class TemplateModelTests(unittest.TestCase):
    def test_json_round_trip_preserves_template(self) -> None:
        template = LabelTemplate(
            name="Shelf label",
            width_mm=100.0,
            height_mm=50.0,
            inner_margins_mm=Margins(2.0, 2.0, 2.0, 2.0),
            outer_margins_mm=Margins(1.0, 1.0, 1.0, 1.0),
        )

        restored = LabelTemplate.from_json(template.to_json())

        self.assertEqual(restored, template)

    def test_template_rejects_duplicate_item_ids(self) -> None:
        template = LabelTemplate(
            name="Duplicate",
            width_mm=20,
            height_mm=20,
            inner_margins_mm=Margins(0, 0, 0, 0),
            outer_margins_mm=Margins(0, 0, 0, 0),
            items=[{"id": "item_1"}, {"id": "item_1"}],
        )

        self.assertTrue(any("unique" in error for error in template.validate()))

    def test_template_json_does_not_mutate_input_object(self) -> None:
        payload = {
            "schema_version": 1,
            "name": "Immutable",
            "width_mm": 20,
            "height_mm": 20,
            "inner_margins_mm": {"top": 0, "bottom": 0, "left": 0, "right": 0},
            "outer_margins_mm": {"top": 0, "bottom": 0, "left": 0, "right": 0},
            "items": [],
        }

        LabelTemplate.from_json(json.dumps(payload))

        self.assertIsInstance(payload["inner_margins_mm"], dict)


class PricingTests(unittest.TestCase):
    def test_case_article_exposes_binding_context(self) -> None:
        article = BulkCaseArticle(
            ident=ProductIdent("A-1", "123456789012", "Coffee"),
            pricing=PricingInfo(12.0),
        )

        self.assertEqual(article.unit_price_inside_case(), 12.0)
        self.assertEqual(article.to_binding_context()["EFFECTIVE_PRICE"], 12.0)

    def test_specialized_binding_context_uses_article_currency(self) -> None:
        article = WeightedArticle(
            ident=ProductIdent("W-1", "123456789012", "Cheese"),
            pricing=PricingInfo(8.5),
            currency="CHF",
        )

        self.assertIn("CHF", article.to_binding_context()["PRIX_KILO"])

    def test_excluded_tax_keeps_ht_price_and_adds_tax(self) -> None:
        article = BulkCaseArticle(
            ident=ProductIdent("A-2", "123456789012", "Tea"),
            pricing=PricingInfo(100.0),
            tax=TaxInfo(tax_rate=20.0, tax_type="EXCLUDED"),
        )

        context = article.to_binding_context()

        self.assertEqual(context["PRICE_HT"], 100.0)
        self.assertEqual(context["TAX_AMOUNT"], 20.0)
        self.assertEqual(context["PRICE_TTC"], 120.0)


class TierResolverTests(unittest.TestCase):
    def test_missing_tier_falls_back_to_base_price(self) -> None:
        result = TierResolver.resolve(
            TierTargetSpec(primary_index=1),
            {"TIERS": [], "SELLING_PRICE": 12.0, "CURRENCY": "EUR"},
        )

        self.assertEqual(result["unit_price"], 12.0)
        self.assertTrue(result["is_fallback"])

    def test_malformed_tiers_fall_back_to_valid_base_price(self) -> None:
        result = TierResolver.resolve(
            TierTargetSpec(primary_index=0),
            {
                "TIERS": [{"qty": 10}, {"qty": 20, "unit_price": "invalid"}],
                "SELLING_PRICE": 15,
            },
        )

        self.assertEqual(result["unit_price"], 15.0)
        self.assertTrue(result["is_fallback"])

    def test_invalid_tier_container_does_not_raise(self) -> None:
        result = TierResolver.resolve(
            TierTargetSpec(primary_index=0, fallback_to_base_price=False),
            {"TIERS": {"qty": 10, "unit_price": 5}},
        )

        self.assertEqual(result, {})


class ApplicationSettingsTests(unittest.TestCase):
    def test_defaults_match_two_route_workflow(self) -> None:
        settings = AppSettings()

        self.assertEqual(settings.generation_edit_policy, GenerationEditPolicy.FREEZE_SNAPSHOT)
        self.assertEqual(settings.binding_object_mode, BindingObjectMode.SPECIALIZED)
        self.assertTrue(settings.confirm_tier_grouping)

    def test_settings_round_trip_preserves_enum_policies(self) -> None:
        restored = AppSettings.from_dict(AppSettings().to_dict())

        self.assertEqual(restored, AppSettings())


class WorkflowStateTests(unittest.TestCase):
    def test_generation_uses_an_immutable_template_snapshot(self) -> None:
        template = LabelTemplate(
            name="Original",
            width_mm=20,
            height_mm=20,
            inner_margins_mm=Margins(0, 0, 0, 0),
            outer_margins_mm=Margins(0, 0, 0, 0),
        )
        snapshot = TemplateSnapshot.from_template(template)
        session = GenerationSession(snapshot)
        template.name = "Editor changed"

        self.assertEqual(session.template_snapshot.load().name, "Original")

    def test_library_indexes_valid_and_invalid_templates(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            valid_path = Path(directory) / "valid.json"
            invalid_path = Path(directory) / "invalid.json"
            valid_path.write_text(
                LabelTemplate(
                    name="Shelf",
                    width_mm=20,
                    height_mm=20,
                    inner_margins_mm=Margins(0, 0, 0, 0),
                    outer_margins_mm=Margins(0, 0, 0, 0),
                ).to_json(),
                encoding="utf-8",
            )
            invalid_path.write_text("not json", encoding="utf-8")

            summaries = TemplateLibrary(directory).discover()

        self.assertEqual(len(summaries), 2)
        self.assertEqual([summary.valid for summary in summaries], [False, True])

    def test_domain_aliases_resolve_to_canonical_binding_keys(self) -> None:
        self.assertEqual(canonical_domain_key("DIV.NAME"), DomainFieldKey.DIV_NAME.value)
        self.assertEqual(canonical_domain_key("EAN"), DomainFieldKey.PRODUCT_SCAN.value)


class TierGroupingTests(unittest.TestCase):
    def test_quantity_one_becomes_base_and_other_rows_become_sorted_tiers(self) -> None:
        rows = [
            ImportedRow("Products", 4, {"SKU": "A", "Purchase Qty": "10", "Price": 8}, {}),
            ImportedRow("Products", 2, {"SKU": "A", "Purchase Qty": "up to 1", "Price": 12}, {}),
            ImportedRow("Products", 3, {"SKU": "A", "Purchase Qty": "5", "Price": 10}, {}),
        ]
        for row in rows:
            row.raw_values.update(row.values)

        result = group_article_rows(rows, TierGroupingConfig(article_key="SKU"))

        self.assertEqual(result[0].base_price, 12.0)
        self.assertEqual([tier["qty"] for tier in result[0].tiers], [5, 10])
        self.assertTrue(result[0].valid)

    def test_conflicting_prices_are_reported_by_default(self) -> None:
        rows = [
            ImportedRow("Products", 2, {"SKU": "A", "QTY": 1, "PRICE": 10}, {}),
            ImportedRow("Products", 3, {"SKU": "A", "QTY": 1, "PRICE": 11}, {}),
        ]
        for row in rows:
            row.raw_values.update(row.values)

        result = group_article_rows(rows, TierGroupingConfig(article_key="SKU"))

        self.assertFalse(result[0].valid)
        self.assertEqual(result[0].issues[0].code, "tier_price_conflict")


class ImpositionTests(unittest.TestCase):
    def test_a4_imposition_centers_labels(self) -> None:
        result = ImpositionCalculator.calculate(100.0, 50.0, Margins(2.0, 2.0, 2.0, 2.0), gap_mm=5.0)

        self.assertEqual(result.total_per_page, result.cols * result.rows)
        self.assertGreaterEqual(result.horizontal_offset_mm, 0.0)
        self.assertGreaterEqual(result.vertical_offset_mm, 0.0)


@unittest.skipIf(Workbook is None, "openpyxl is not installed")
class ExcelImportTests(unittest.TestCase):
    def _workbook_path(self) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Products"
        sheet.append(["EAN", "Name", "Price", "Date", "TIER_1_QTY", "TIER_1_PRICE", "TIER_1_UNIT"])
        sheet.append(["0012345678905", "Coffee", "12,50", date(2026, 9, 5), 10, 11.5, "PCE"])
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        workbook.save(handle.name)
        return Path(handle.name)

    def test_xlsx_import_preserves_identifier_and_normalizes_types(self) -> None:
        path = self._workbook_path()
        try:
            table = read_xlsx(path, "Products")
            profile = BindingProfile(mappings=[
                ColumnMapping("EAN", "PRODUCT_SCAN", "barcode", required=True),
                ColumnMapping("Name", "ITEMNAME", "text", required=True),
                ColumnMapping("Price", "SELLING_PRICE", "currency"),
                ColumnMapping("Date", "DATE", "date"),
            ])
            result = normalize_row(table.rows[0], profile)

            self.assertTrue(result.valid)
            self.assertEqual(result.record["PRODUCT_SCAN"], "0012345678905")
            self.assertEqual(result.record["SELLING_PRICE"], 12.5)
            self.assertEqual(result.record["DATE"], "2026-09-05")
            self.assertEqual(result.record["TIERS"][0]["qty"], 10)
            self.assertEqual(result.record["TIERS"][0]["unit_price"], 11.5)
        finally:
            path.unlink(missing_ok=True)

    def test_duplicate_headers_are_reported_and_made_stable(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Name", "Name", None])
        sheet.append(["A", "B", "C"])
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        path = Path(handle.name)
        try:
            workbook.save(path)
            table = read_xlsx(path)
            self.assertEqual(table.headers, ["Name", "Name_2", "COLUMN_3"])
            self.assertTrue(any(issue.code == "duplicate_header" for issue in table.issues))
            self.assertTrue(any(issue.code == "blank_header" for issue in table.issues))
        finally:
            path.unlink(missing_ok=True)

    def test_required_and_invalid_tier_values_are_reported(self) -> None:
        row = ImportedRow("Products", 4, {"Code": "", "TIER_1_QTY": 10, "TIER_1_PRICE": "bad"}, {"Code": "", "TIER_1_QTY": 10, "TIER_1_PRICE": "bad"})
        profile = BindingProfile(mappings=[ColumnMapping("Code", "PRODUCT_SCAN", "barcode", required=True)])

        result = normalize_row(row, profile)

        self.assertFalse(result.valid)
        self.assertEqual({issue.code for issue in result.issues}, {"conversion_error", "invalid_tier"})

    def test_batch_normalization_isolates_bad_rows(self) -> None:
        rows = [
            ImportedRow("Products", 2, {"Price": "10"}, {"Price": "10"}),
            ImportedRow("Products", 3, {"Price": "bad"}, {"Price": "bad"}),
        ]
        table = type("Table", (), {"headers": ["Price"], "iter_rows": lambda self: iter(rows)})()
        profile = BindingProfile(mappings=[ColumnMapping("Price", "SELLING_PRICE", "number", required=True)])

        report = normalize_workbook(table, profile, "skip_invalid_rows")

        self.assertEqual(report.succeeded, 1)
        self.assertEqual(report.failed, 1)
        self.assertIsNone(report.results[1].record)

    def test_batch_png_export_uses_stable_row_names(self) -> None:
        rows = [ImportedRow("Products", 2, {"Name": "A"}, {"Name": "A"})]
        table = type("Table", (), {"headers": ["Name"], "iter_rows": lambda self: iter(rows)})()
        profile = BindingProfile(mappings=[ColumnMapping("Name", "ITEMNAME", "text", required=True)])
        report = normalize_workbook(table, profile)

        class FakeCanvas:
            def export_png(self, path, dpi, data_record):
                Path(path).write_bytes(b"PNG")

        with tempfile.TemporaryDirectory() as directory:
            exported = export_rows_as_png(FakeCanvas(), report, directory)
            self.assertEqual(exported.succeeded, 1)
            self.assertTrue((Path(directory) / "label_00002.png").exists())

    def test_batch_normalization_can_be_cancelled(self) -> None:
        rows = [
            ImportedRow("Products", 2, {"Name": "A"}, {"Name": "A"}),
            ImportedRow("Products", 3, {"Name": "B"}, {"Name": "B"}),
        ]
        table = type("Table", (), {"headers": ["Name"], "iter_rows": lambda self: iter(rows)})()
        profile = BindingProfile(mappings=[ColumnMapping("Name", "ITEMNAME", "text")])

        report = normalize_workbook(table, profile, progress=lambda current, total: current == 0)

        self.assertEqual(len(report.results), 1)
        self.assertEqual(report.results[0].row_number, 2)


if __name__ == "__main__":
    unittest.main()
