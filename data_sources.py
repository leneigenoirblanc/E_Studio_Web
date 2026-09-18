"""Typed workbook access for spreadsheet-backed label data."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - dependency installation concern
    load_workbook = None


@dataclass(frozen=True)
class ImportIssue:
    """A source or row issue that can be shown without exposing a traceback."""

    code: str
    message: str
    severity: str = "error"
    sheet: str | None = None
    row_number: int | None = None
    column: str | None = None


@dataclass(frozen=True)
class ImportedRow:
    sheet: str
    row_number: int
    values: dict[str, Any]
    raw_values: dict[str, Any]


@dataclass
class WorkbookTable:
    """A selected worksheet with stable, de-duplicated headers."""

    path: str
    sheet: str
    headers: list[str]
    original_headers: dict[str, str]
    rows: list[ImportedRow] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)

    def iter_rows(self) -> Iterator[ImportedRow]:
        yield from self.rows


class WorkbookReadError(ValueError):
    """Raised when a workbook cannot be read or has no usable table."""


def _stable_headers(values: tuple[Any, ...]) -> tuple[list[str], dict[str, str], list[ImportIssue]]:
    headers: list[str] = []
    original_headers: dict[str, str] = {}
    issues: list[ImportIssue] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, 1):
        original = "" if value is None else str(value).strip()
        base = original or f"COLUMN_{index}"
        if not original:
            issues.append(ImportIssue("blank_header", f"Column {index} has no header.", "warning", column=base))
        count = seen.get(base, 0) + 1
        seen[base] = count
        header = base if count == 1 else f"{base}_{count}"
        if count > 1:
            issues.append(ImportIssue("duplicate_header", f"Duplicate header {base!r}; imported as {header!r}.", "warning", column=base))
        headers.append(header)
        original_headers[header] = original
    return headers, original_headers, issues


def read_xlsx(path: str | Path, sheet_name: str | None = None, header_row: int = 1) -> WorkbookTable:
    """Read one worksheet from an xlsx/xlsm workbook while preserving cell types."""
    if load_workbook is None:
        raise WorkbookReadError("openpyxl is required to import Excel workbooks.")
    source = Path(path)
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise WorkbookReadError("Only .xlsx and .xlsm files are supported. Convert legacy .xls/.xlsb files first.")
    if header_row < 1:
        raise WorkbookReadError("The header row must be at least 1.")
    try:
        workbook = load_workbook(
            source,
            read_only=True,
            data_only=True,
            keep_vba=source.suffix.lower() == ".xlsm",
        )
    except FileNotFoundError as error:
        raise WorkbookReadError(f"Workbook not found: {source}") from error
    except (OSError, ValueError, KeyError) as error:
        raise WorkbookReadError(f"Unable to read workbook {source.name}: {error}") from error

    names = workbook.sheetnames
    if not names:
        raise WorkbookReadError("The workbook contains no worksheets.")
    selected = sheet_name or names[0]
    if selected not in names:
        raise WorkbookReadError(f"Worksheet {selected!r} was not found. Available sheets: {', '.join(names)}")
    worksheet = workbook[selected]
    iterator = worksheet.iter_rows(values_only=False)
    headers: list[str] | None = None
    original_headers: dict[str, str] = {}
    issues: list[ImportIssue] = []
    rows: list[ImportedRow] = []

    for row_number, cells in enumerate(iterator, 1):
        if row_number < header_row:
            continue
        values = tuple(cell.value for cell in cells)
        if row_number == header_row:
            headers, original_headers, issues = _stable_headers(values)
            continue
        if headers is None:
            continue
        raw_values = {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(headers)
        }
        if all(value is None for value in raw_values.values()):
            continue
        rows.append(ImportedRow(selected, row_number, dict(raw_values), raw_values))

    workbook.close()
    if headers is None:
        raise WorkbookReadError(f"Worksheet {selected!r} has no header row at {header_row}.")
    return WorkbookTable(str(source), selected, headers, original_headers, rows, issues)
